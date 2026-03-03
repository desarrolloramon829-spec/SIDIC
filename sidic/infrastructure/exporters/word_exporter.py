"""
Exportador a Microsoft Word (.docx).

Genera documentos Word profesionales con tablas estilizadas,
gráficos embebidos, encabezados policiales y cuadros de referencia.
"""
from __future__ import annotations

import io
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd

from sidic.infrastructure.exporters.base_exporter import (
    BaseExporter,
    ExportOptions,
    ExportResult,
)

logger = logging.getLogger(__name__)


class WordExporter(BaseExporter):
    """
    Exportador de reportes a Word (.docx).

    Genera documento profesional con estilos policiales
    (rojo institucional, tablas con bordes, gráficos).
    """

    def export(
        self,
        output_path: str | Path,
        tables: Dict[str, pd.DataFrame],
        charts: Optional[Dict[str, bytes]] = None,
        options: Optional[ExportOptions] = None,
    ) -> ExportResult:
        try:
            from docx import Document
            from docx.shared import Cm
        except ImportError as e:
            msg = f"python-docx no instalado: {e}"
            self._last_error = msg
            return ExportResult(success=False, error=msg)

        options = options or ExportOptions()
        charts = charts or {}
        output_path = Path(output_path)
        self._ensure_dir(output_path)

        doc = Document()
        _setup_margins(doc, cm=2)

        try:
            # Título
            _add_heading(doc, options.titulo, level=1)
            if options.subtitulo:
                _add_paragraph(doc, options.subtitulo, bold=True, align="center")

            # Info general
            info = [
                ("Jurisdicción", options.jurisdiccion),
                ("Dependencia", options.dependencia),
                ("Autor", options.autor),
            ]
            for label, val in info:
                if val:
                    _add_paragraph(doc, f"{label}: {val}", bold=False)

            # Secciones de tablas
            _SECTIONS = [
                ("cuadro_referencia", "CUADRO DE REFERENCIA", True),
                ("delitos", "DELITOS CON MODALIDADES", True),
                ("dias_semana", "DÍAS DE LA SEMANA", True),
                ("franja_horaria", "FRANJA HORARIA", True),
                ("movilidad", "MEDIOS DE MOVILIDAD", True),
                ("armas", "ARMAS / MEDIOS EN ROBOS AGRAVADOS", True),
                ("ambito", "ÁMBITO DE OCURRENCIA", True),
                ("esclarecimiento", "ÍNDICE DE ESCLARECIMIENTO", True),
                ("matriz_delito_dia", "DELITOS POR DÍA DE LA SEMANA", options.incluir_matrices),
                ("matriz_delito_franja", "DELITOS POR FRANJA HORARIA", options.incluir_matrices),
                ("mencionados", "MENCIONADOS", options.incluir_mencionados),
                ("aprehendidos", "APREHENDIDOS", options.incluir_aprehendidos),
                ("aprehendidos_clasificacion", "CLASIFICACIÓN DE APREHENDIDOS", options.incluir_aprehendidos),
                ("comparativa_detallada", "CUADRO COMPARATIVO", options.incluir_comparativos),
                ("comparativa_general", "CUADRO COMPARATIVO GENERAL", options.incluir_comparativos),
            ]

            for key, title, enabled in _SECTIONS:
                if not enabled:
                    continue
                if key not in tables or tables[key].empty:
                    continue
                # Evitar duplicar comparativa
                if key == "comparativa_general" and "comparativa_detallada" in tables:
                    continue

                doc.add_page_break()
                _add_heading(doc, title, level=2)
                _add_table(doc, tables[key])

                # Gráfico
                if options.incluir_graficos and key in charts:
                    _add_image(doc, charts[key])

            # Guardar
            return self._save_with_retry(lambda p: doc.save(p), output_path)

        except Exception as e:
            msg = f"Error exportando Word: {e}"
            self._last_error = msg
            logger.exception(msg)
            return ExportResult(success=False, error=msg)


# ═══════════════════════════════════════════════════════════════════════════
# FUNCIONES INTERNAS
# ═══════════════════════════════════════════════════════════════════════════


def _setup_margins(doc, cm: float = 2):
    from docx.shared import Cm as CmUnit

    for section in doc.sections:
        section.top_margin = CmUnit(cm)
        section.bottom_margin = CmUnit(cm)
        section.left_margin = CmUnit(cm)
        section.right_margin = CmUnit(cm)


def _add_heading(doc, text: str, level: int = 1):
    from docx.shared import RGBColor

    heading = doc.add_heading(text, level=level)
    if level <= 2:
        for run in heading.runs:
            run.font.color.rgb = RGBColor(204, 0, 0)


def _add_paragraph(doc, text: str, bold: bool = False, align: str = "left"):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    mapping = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }
    p.alignment = mapping.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    return p


def _set_cell_shading(cell, color: str):
    """Aplica color de fondo a celda Word."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _add_table(doc, df: pd.DataFrame):
    """Agrega una tabla Word desde DataFrame."""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Pt, RGBColor

    if df.empty:
        return

    # Filtrar columnas internas
    display_cols = [c for c in df.columns if c not in ("color_fila", "tendencia", "Color")]

    n_rows = len(df) + 1
    n_cols = len(display_cols)

    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Headers
    for ci, col_name in enumerate(display_cols):
        cell = table.rows[0].cells[ci]
        cell.text = str(col_name)
        _set_cell_shading(cell, "FF0000")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)

    # Data rows
    for ri, (_, row) in enumerate(df.iterrows()):
        first_val = str(row.iloc[0]).upper() if len(row) > 0 else ""
        is_total = "TOTAL" in first_val

        for ci, col_name in enumerate(display_cols):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(row[col_name]) if pd.notna(row[col_name]) else ""

            if is_total:
                _set_cell_shading(cell, "FFFF00")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.size = Pt(10)
            elif ri % 2 == 0:
                _set_cell_shading(cell, "F0F0F0")


def _add_image(doc, image_bytes: bytes, width_inches: float = 5.5):
    """Inserta imagen PNG en el documento."""
    from docx.shared import Inches

    try:
        doc.add_picture(io.BytesIO(image_bytes), width=Inches(width_inches))
    except Exception as e:
        logger.warning(f"No se pudo insertar imagen: {e}")
