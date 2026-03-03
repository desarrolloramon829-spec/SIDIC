"""
Exportador a Microsoft Excel (.xlsx).

Genera archivos Excel con tablas estilizadas, gráficos embebidos y
cuadro de referencia con símbolos coloreados.
"""
from __future__ import annotations

import io
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

from sidic.config import get_section
from sidic.infrastructure.exporters.base_exporter import (
    BaseExporter,
    ExportOptions,
    ExportResult,
)

logger = logging.getLogger(__name__)


class ExcelExporter(BaseExporter):
    """
    Exportador de reportes a Excel con estilos policiales.

    Colores por defecto:
        - Encabezados: rojo policial + texto blanco
        - Totales: amarillo + texto negro
        - Subtotales: azul + texto blanco
        - Variación positiva: rojo (delito subió)
        - Variación negativa: azul (delito bajó)
    """

    def export(
        self,
        output_path: str | Path,
        tables: Dict[str, pd.DataFrame],
        charts: Optional[Dict[str, bytes]] = None,
        options: Optional[ExportOptions] = None,
    ) -> ExportResult:
        from openpyxl import Workbook

        options = options or ExportOptions()
        charts = charts or {}
        output_path = Path(output_path)
        self._ensure_dir(output_path)

        colors = _load_colors()
        wb = Workbook()

        try:
            # 1. Hoja resumen
            _write_resumen(wb, options)

            # 2. Cuadro de referencia
            if options.incluir_cuadro_referencia and "cuadro_referencia" in tables:
                _write_cuadro_referencia(wb, tables["cuadro_referencia"], colors)

            # 3. Tablas regulares
            _TABLE_MAP = [
                ("delitos", "Delitos", "DELITOS CON MODALIDADES"),
                ("dias_semana", "Días Semana", "DÍAS DE LA SEMANA"),
                ("franja_horaria", "Franja Horaria", "FRANJA HORARIA"),
                ("movilidad", "Movilidad", "MEDIOS DE MOVILIDAD"),
                ("armas", "Armas", "ARMAS/MEDIOS EN ROBOS AGRAVADOS"),
                ("ambito", "Ámbito", "ÁMBITO DE OCURRENCIA"),
                ("esclarecimiento", "Esclarecimiento", "ÍNDICE DE ESCLARECIMIENTO"),
            ]
            for key, sheet_name, title in _TABLE_MAP:
                if key in tables and not tables[key].empty:
                    _write_table_sheet(
                        wb,
                        self._safe_sheet_name(sheet_name),
                        title,
                        tables[key],
                        charts.get(key) if options.incluir_graficos else None,
                        colors,
                    )

            # 4. Matrices cruzadas
            if options.incluir_matrices:
                for key, name, title in [
                    ("matriz_delito_dia", "Delitos x Día", "DELITOS POR DÍA DE LA SEMANA"),
                    ("matriz_delito_franja", "Delitos x Franja", "DELITOS POR FRANJA HORARIA"),
                ]:
                    if key in tables and not tables[key].empty:
                        _write_table_sheet(
                            wb,
                            self._safe_sheet_name(name),
                            title,
                            tables[key],
                            charts.get(key) if options.incluir_graficos else None,
                            colors,
                        )

            # 5. Mencionados / Aprehendidos
            if options.incluir_mencionados and "mencionados" in tables:
                _write_table_sheet(
                    wb, "Mencionados", "MENCIONADOS", tables["mencionados"], None, colors
                )
            if options.incluir_aprehendidos and "aprehendidos" in tables:
                _write_table_sheet(
                    wb, "Aprehendidos", "APREHENDIDOS", tables["aprehendidos"], None, colors
                )
            if "aprehendidos_clasificacion" in tables:
                _write_table_sheet(
                    wb,
                    "Clasif Aprehendidos",
                    "CLASIFICACIÓN DE APREHENDIDOS",
                    tables["aprehendidos_clasificacion"],
                    None,
                    colors,
                )

            # 6. Comparativa
            if options.incluir_comparativos:
                if "comparativa_detallada" in tables and not tables["comparativa_detallada"].empty:
                    _write_comparative_sheet(
                        wb,
                        tables["comparativa_detallada"],
                        charts.get("comparativa_delitos") if options.incluir_graficos else None,
                        colors,
                    )
                elif "comparativa_general" in tables:
                    _write_table_sheet(
                        wb,
                        "Comparativa",
                        "CUADRO COMPARATIVO",
                        tables["comparativa_general"],
                        None,
                        colors,
                    )

            # Guardar
            return self._save_with_retry(lambda p: wb.save(p), output_path)

        except ImportError as e:
            msg = f"Dependencia faltante: {e}"
            self._last_error = msg
            return ExportResult(success=False, error=msg)
        except Exception as e:
            msg = f"Error exportando Excel: {e}"
            self._last_error = msg
            logger.exception(msg)
            return ExportResult(success=False, error=msg)


# ═══════════════════════════════════════════════════════════════════════════
# FUNCIONES INTERNAS DE ESCRITURA
# ═══════════════════════════════════════════════════════════════════════════

def _load_colors() -> Dict[str, str]:
    """Carga colores desde config o usa defaults."""
    try:
        return get_section("colores_excel")
    except Exception:
        return {
            "header_bg": "FF0000",
            "header_fg": "FFFFFF",
            "total_bg": "FFFF00",
            "total_fg": "000000",
            "subtotal_bg": "0000FF",
            "subtotal_fg": "FFFFFF",
            "up_bg": "FF0000",
            "down_bg": "0000FF",
            "alt_row": "F0F0F0",
        }


def _styled_cell(ws, row, col, value, *, font_kw=None, fill_color=None, align="center"):
    """Escribe una celda con estilo."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    cell = ws.cell(row=row, column=col, value=value)
    cell.border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if fill_color:
        cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    if font_kw:
        cell.font = Font(**font_kw)
    return cell


def _write_resumen(wb, options: ExportOptions):
    """Escribe la hoja de resumen."""
    from openpyxl.styles import Alignment, Font

    ws = wb.active
    ws.title = "Resumen"

    ws.cell(row=1, column=1, value=options.titulo).font = Font(bold=True, size=14, color="CC0000")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)

    info_rows = [
        ("Jurisdicción:", options.jurisdiccion or "No especificada"),
        ("Dependencia:", options.dependencia or ""),
        ("Autor:", options.autor or ""),
    ]
    row = 3
    for label, val in info_rows:
        if val:
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row, column=2, value=val)
            row += 1

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 45


def _write_table_sheet(
    wb,
    sheet_name: str,
    title: str,
    df: pd.DataFrame,
    chart_bytes: Optional[bytes],
    colors: Dict[str, str],
):
    """Escribe una hoja con título, tabla y gráfico opcional."""
    from openpyxl.styles import Font

    ws = wb.create_sheet(title=sheet_name[:31])
    n_cols = len(df.columns)

    # Título
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14, color="CC0000")
    if n_cols > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)

    # Tabla
    final_row = _write_dataframe(ws, df, start_row=3, colors=colors)

    # Gráfico
    if chart_bytes:
        _insert_image(ws, chart_bytes, f"A{final_row + 2}")


def _write_dataframe(
    ws,
    df: pd.DataFrame,
    start_row: int = 1,
    start_col: int = 1,
    colors: Optional[Dict[str, str]] = None,
) -> int:
    """Escribe DataFrame con estilos, retorna la fila siguiente."""
    colors = colors or {}
    header_bg = colors.get("header_bg", "FF0000")
    header_fg = colors.get("header_fg", "FFFFFF")
    total_bg = colors.get("total_bg", "FFFF00")
    total_fg = colors.get("total_fg", "000000")
    alt_row_bg = colors.get("alt_row", "F0F0F0")

    n_cols = len(df.columns)
    cur = start_row

    # Headers
    for ci, col_name in enumerate(df.columns, start=start_col):
        _styled_cell(
            ws, cur, ci, col_name,
            font_kw={"bold": True, "color": header_fg, "size": 11},
            fill_color=header_bg,
        )
    cur += 1

    # Data
    for ri, (_, row) in enumerate(df.iterrows()):
        first_val = str(row.iloc[0]).upper() if len(row) > 0 else ""
        is_total = "TOTAL" in first_val
        is_subtotal = "SUBTOTAL" in first_val

        for ci, val in enumerate(row, start=start_col):
            if is_total:
                _styled_cell(
                    ws, cur, ci, val,
                    font_kw={"bold": True, "color": total_fg, "size": 11},
                    fill_color=total_bg,
                )
            elif is_subtotal:
                _styled_cell(
                    ws, cur, ci, val,
                    font_kw={"bold": True, "color": "FFFFFF", "size": 10},
                    fill_color=colors.get("subtotal_bg", "0000FF"),
                )
            else:
                bg = alt_row_bg if ri % 2 == 0 else "FFFFFF"
                align = "left" if ci == start_col else "center"
                _styled_cell(
                    ws, cur, ci, val,
                    font_kw={"size": 10},
                    fill_color=bg,
                    align=align,
                )
        cur += 1

    # Auto-width
    from openpyxl.utils import get_column_letter

    for ci, col_name in enumerate(df.columns, start=start_col):
        max_len = max(
            len(str(col_name)),
            *(len(str(v)) for v in df[col_name]),
            0,
        )
        ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 3, 50)

    return cur


def _write_cuadro_referencia(wb, df: pd.DataFrame, colors: Dict[str, str]):
    """Escribe cuadro de referencia con símbolos coloreados."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    ws = wb.create_sheet(title="Cuadro Referencia")
    ws.cell(row=1, column=1, value="CUADRO DE REFERENCIA").font = Font(
        bold=True, size=14, color="CC0000"
    )
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)

    headers = [c for c in df.columns if c != "Color"]
    cur = 3
    header_bg = colors.get("header_bg", "FF0000")
    for ci, h in enumerate(headers, start=1):
        _styled_cell(
            ws, cur, ci, h,
            font_kw={"bold": True, "color": "FFFFFF", "size": 11},
            fill_color=header_bg,
        )
    cur += 1

    for _, row in df.iterrows():
        color_hex = str(row.get("Color", "")).replace("#", "")
        for ci, col_name in enumerate(headers, start=1):
            val = row[col_name]
            cell = _styled_cell(ws, cur, ci, val, font_kw={"size": 10})
            if ci == 1 and color_hex:
                # Símbolo grande y coloreado
                if color_hex.upper() in ("FFFFFF", "FFFF00"):
                    cell.fill = PatternFill(
                        start_color="000000", end_color="000000", fill_type="solid"
                    )
                cell.font = Font(bold=True, color=color_hex, size=14)
                cell.alignment = Alignment(horizontal="center", vertical="center")
        cur += 1

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 12


def _write_comparative_sheet(
    wb,
    df: pd.DataFrame,
    chart_bytes: Optional[bytes],
    colors: Dict[str, str],
):
    """Escribe hoja comparativa con colores según tendencia."""
    from openpyxl.styles import Font

    ws = wb.create_sheet(title="Comparativa")
    ws.cell(row=1, column=1, value="CUADRO COMPARATIVO ENTRE PERÍODOS").font = Font(
        bold=True, size=14, color="CC0000"
    )

    display_cols = [c for c in df.columns if c not in ("color_fila", "tendencia")]
    n_cols = len(display_cols)

    header_bg = colors.get("header_bg", "FF0000")
    total_bg = colors.get("total_bg", "FFFF00")
    up_bg = colors.get("up_bg", "FF0000")
    down_bg = colors.get("down_bg", "0000FF")

    cur = 3
    for ci, col in enumerate(display_cols, start=1):
        _styled_cell(
            ws, cur, ci, col,
            font_kw={"bold": True, "color": "FFFFFF", "size": 10},
            fill_color=header_bg,
        )
    ws.row_dimensions[cur].height = 40
    cur += 1

    for _, row in df.iterrows():
        tendencia = row.get("tendencia", "") if "tendencia" in row.index else ""
        color_fila = (
            str(row.get("color_fila", "#FFFFFF")).replace("#", "")
            if "color_fila" in row.index
            else "FFFFFF"
        )
        is_total = tendencia == "total" or "TOTAL" in str(row.iloc[0]).upper()

        for ci, col_name in enumerate(display_cols, start=1):
            val = row[col_name]
            if is_total:
                _styled_cell(
                    ws, cur, ci, val,
                    font_kw={"bold": True, "color": "000000", "size": 11},
                    fill_color=total_bg,
                )
            elif color_fila == up_bg:
                _styled_cell(
                    ws, cur, ci, val,
                    font_kw={"color": "FFFFFF", "size": 10},
                    fill_color=up_bg,
                )
            elif color_fila == down_bg:
                _styled_cell(
                    ws, cur, ci, val,
                    font_kw={"color": "FFFFFF", "size": 10},
                    fill_color=down_bg,
                )
            else:
                align = "left" if ci == 1 else "center"
                _styled_cell(ws, cur, ci, val, font_kw={"size": 10}, align=align)
        cur += 1

    from openpyxl.utils import get_column_letter

    widths = {1: 40, 2: 18, 3: 18, 4: 35, 5: 35}
    for ci, w in widths.items():
        if ci <= n_cols:
            ws.column_dimensions[get_column_letter(ci)].width = w

    if chart_bytes:
        _insert_image(ws, chart_bytes, f"A{cur + 2}")


def _insert_image(ws, image_bytes: bytes, cell: str):
    """Inserta imagen PNG en la hoja."""
    try:
        from openpyxl.drawing.image import Image as XLImage

        img = XLImage(io.BytesIO(image_bytes))
        img.width = 600
        img.height = 400
        ws.add_image(img, cell)
    except Exception as e:
        logger.warning(f"No se pudo insertar imagen: {e}")
